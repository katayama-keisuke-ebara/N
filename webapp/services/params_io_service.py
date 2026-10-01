from __future__ import annotations

import io
import zipfile

import openpyxl
from openpyxl.styles import Font
from pydantic import ValidationError

from webapp.schemas import PartsParamsPayload

MAX_FILE_BYTES = 5 * 1024 * 1024
_SETTINGS = {
    "定尺幅(mm)": "sheet_w", "定尺高さ(mm)": "sheet_h", "材料厚み(mm)": "thickness",
    "クリアランス(mm)": "clearance", "密度(g/cm³)": "density",
}
_HEADERS = ("種別", "ファイル名", "幅(mm)", "高さ(mm)", "円単位", "円の値(mm)", "数量")


class ParamsImportError(ValueError):
    pass


class ParamsIOService:
    def build_workbook(self, data: dict) -> bytes:
        params = PartsParamsPayload.model_validate(data)
        workbook = openpyxl.Workbook()
        settings = workbook.active
        settings.title = "設定"
        settings.append(("項目", "値"))
        for label, key in _SETTINGS.items():
            settings.append((label, getattr(params, key)))
        parts = workbook.create_sheet("部品リスト")
        parts.append(("STEP はファイル名のみ保存します。読み込み後に元の STEP ファイルを再選択してください。",))
        parts.merge_cells("A1:G1")
        parts.append(_HEADERS)
        for part in params.step_parts:
            parts.append(("STEP", part.filename, None, None, None, None, part.quantity))
            parts.cell(parts.max_row, 2).data_type = "s"
        for part in params.manual_parts:
            if part.shape_type == "rectangle":
                parts.append(("矩形", None, part.width, part.height, None, None, part.quantity))
            else:
                unit = "直径" if part.circle_unit == "diameter" else "半径"
                parts.append(("円形", None, None, None, unit, part.circle_value, part.quantity))
        for sheet, header_row in ((settings, 1), (parts, 2)):
            sheet.freeze_panes = f"A{header_row + 1}"
            for cell in sheet[header_row]:
                cell.font = Font(bold=True)
            for column in ("A", "B", "C", "D", "E", "F", "G"):
                sheet.column_dimensions[column].width = 24 if column != "B" else 48
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def parse_workbook(self, content: bytes) -> dict:
        workbook = None
        try:
            if len(content) > MAX_FILE_BYTES:
                raise ParamsImportError("Excel ファイルは 5 MB 以下にしてください")
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ParamsImportError("Excel ファイルの展開サイズが大きすぎます")
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            if not {"設定", "部品リスト"}.issubset(workbook.sheetnames):
                raise ParamsImportError("「設定」と「部品リスト」シートが必要です")
            settings, parts = workbook["設定"], workbook["部品リスト"]
            for sheet in (settings, parts):
                if (sheet.max_row or 0) > 20002 or (sheet.max_column or 0) > 7:
                    raise ParamsImportError("Excel の行数または列数が多すぎます")
                for row in sheet.iter_rows():
                    if any(cell.data_type in ("f", "e") for cell in row):
                        raise ParamsImportError("数式やエラー値ではなく値を入力してください")
            if tuple(next(parts.iter_rows(min_row=2, max_row=2, max_col=7, values_only=True))) != _HEADERS:
                raise ParamsImportError("部品リストの見出しが一致しません")
            data: dict = {"step_parts": [], "manual_parts": []}
            for label, value in settings.iter_rows(min_row=2, max_col=2, values_only=True):
                if label in _SETTINGS:
                    key = _SETTINGS[label]
                    if key in data:
                        raise ParamsImportError(f"設定項目が重複しています: {label}")
                    data[key] = value
            for row_number, row in enumerate(parts.iter_rows(min_row=3, max_col=7, values_only=True), 3):
                if all(value is None for value in row):
                    continue
                kind, filename, width, height, unit, value, quantity = row
                quantity = 1 if quantity is None else quantity
                if kind == "STEP":
                    data["step_parts"].append({"filename": filename, "quantity": quantity})
                elif kind == "矩形":
                    data["manual_parts"].append({
                        "shape_type": "rectangle", "width": width, "height": height, "quantity": quantity,
                    })
                elif kind == "円形" and unit in ("直径", "半径"):
                    data["manual_parts"].append({
                        "shape_type": "circle", "circle_unit": "diameter" if unit == "直径" else "radius",
                        "circle_value": value, "quantity": quantity,
                    })
                else:
                    raise ParamsImportError(f"部品リスト {row_number} 行目の種別または円単位が不正です")
            return PartsParamsPayload.model_validate(data).model_dump()
        except ParamsImportError:
            raise
        except ValidationError as exc:
            location = ".".join(str(item) for item in exc.errors()[0]["loc"])
            raise ParamsImportError(f"パラメータが不足または範囲外です: {location}") from exc
        except Exception as exc:
            raise ParamsImportError("パラメータ用の有効な Excel (.xlsx) を選択してください") from exc
        finally:
            if workbook is not None:
                workbook.close()