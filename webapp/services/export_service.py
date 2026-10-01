from __future__ import annotations

import logging
import shutil
from datetime import UTC, datetime
from io import BytesIO
from math import floor
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.export.nesting_render import render_nesting_png

logger = logging.getLogger(__name__)

_SUMMARY_HEADERS = [
    "処理日時", "部品一覧", "分布案", "取り数合計", "歩留まり (%)",
    "残材面積 (mm²)", "穴面積合計 (mm²)", "定尺",
    "密度 (g/cm³)", "部材重量合計 (kg)", "端材重量 (kg)",
    "部材重量 計算式", "端材重量 計算式",
]

_PARTS_HEADERS = [
    "部品名", "数量指定", "配置数", "部品面積 (mm²)", "部品体積 (mm³)", "穴面積合計 (mm²)",
    "重量", "重量 計算式",
]


def _weight_kg(area_mm2: float, thickness_mm: float, density_g_cm3: float) -> float:
    """mm² × mm × g/cm³ を kg に換算する (mm³をcm³化してgを得てからkgへ)。"""
    return area_mm2 * thickness_mm * density_g_cm3 / 1_000_000.0


def _format_weight(weight_kg: float) -> str:
    return f"{round(weight_kg, 4)}kg ({floor(weight_kg * 1000 + 0.5)}g)"


def _weight_formula(area_mm2: float, thickness_mm: float, density_g_cm3: float, count: int | None = None) -> str:
    """重量計算の根拠 (検算用) を人が読める式の文字列として返す。"""
    count_part = f" × {count}個" if count is not None else ""
    return (
        f"(= {area_mm2:.4f}mm²{count_part} × {thickness_mm}mm × {density_g_cm3}g/cm³ ÷ 1,000,000)"
    )


def _scrap_area(candidate: dict) -> float:
    """シート面積から配置部品の正味面積と穴面積を差し引いた端材面積。"""
    return max(
        0.0,
        candidate.get("sheet_area", 0) - candidate.get("total_used_area", 0)
        - candidate.get("total_hole_area", 0),
    )


class ExportService:
    """DXF コピーと Excel 出力を担当するサービス。"""

    def save_dxf(self, src_path: str, dst_path: str) -> None:
        """既存 DXF をユーザー指定のパスへコピーする。"""
        shutil.copy2(src_path, dst_path)
        logger.info("DXF を保存しました: %s", dst_path)

    def save_excel(
        self,
        results: list[dict],
        candidate: dict,
        params: dict,
        dst_path: str,
    ) -> None:
        """計算結果と加工パラメータを Excel ファイル (.xlsx) に書き出す。"""
        wb = openpyxl.Workbook()
        ws = wb.active or wb.create_sheet()
        ws.title = "Summary"

        sheet_size = f"{params.get('sheet_w', 0):.0f}×{params.get('sheet_h', 0):.0f} mm"
        sheet_area = params.get("sheet_w", 0) * params.get("sheet_h", 0)
        count_per_part = candidate.get("count_per_part", {})

        header_font = Font(bold=True)
        left = Alignment(horizontal="left")

        rows: list[tuple] = [
            ("項目", "値", "単位"),
            ("処理日時", datetime.now(tz=UTC).astimezone().strftime("%Y-%m-%d %H:%M:%S"), ""),
            ("部品数", len(results), "種類"),
            ("定尺", sheet_size, ""),
            ("分布案", candidate.get("strategy_name", ""), ""),
            ("取り数合計", sum(count_per_part.values()), "個"),
            ("歩留まり", round(candidate.get("yield_pct", 0), 2), "%"),
            ("残材面積", round(candidate.get("remaining_area", 0), 4), "mm²"),
            ("穴面積合計", round(candidate.get("total_hole_area", 0), 4), "mm²"),
            ("材料厚み", params.get("thickness", 0), "mm"),
            ("クリアランス", params.get("clearance", 0), "mm"),
            ("定尺面積", round(sheet_area, 4), "mm²"),
            ("密度", params.get("density", 0), "g/cm³"),
            (
                "部材重量合計",
                _format_weight(_weight_kg(
                    candidate.get("total_used_area", 0), params.get("thickness", 0),
                    params.get("density", 0),
                )),
                "",
            ),
            (
                "部材重量合計 計算式",
                _weight_formula(
                    candidate.get("total_used_area", 0), params.get("thickness", 0),
                    params.get("density", 0),
                ),
                "",
            ),
            (
                "端材重量",
                _format_weight(_weight_kg(
                    _scrap_area(candidate), params.get("thickness", 0), params.get("density", 0),
                )),
                "",
            ),
            (
                "端材重量 計算式",
                _weight_formula(
                    _scrap_area(candidate), params.get("thickness", 0), params.get("density", 0),
                ),
                "",
            ),
        ]

        for row_idx, row_data in enumerate(rows, start=1):
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = left
                if row_idx == 1:
                    cell.font = header_font
        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 24

        parts_ws = wb.create_sheet("Parts")
        for col_idx, header in enumerate(_PARTS_HEADERS, start=1):
            cell = parts_ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_idx, result in enumerate(results, start=2):
            part_id = result.get("part_id") or f"{Path(result['step_path']).stem}_{row_idx - 2}"
            quantity = result.get("quantity", 1)
            quantity_label = "自動(最大)" if quantity <= 0 else quantity
            count = count_per_part.get(part_id, 0)
            part_area = result.get("flattened_area", 0)
            thickness = params.get("thickness", 0)
            density = params.get("density", 0)
            parts_ws.append([
                result.get("display_name") or Path(result["step_path"]).stem,
                quantity_label,
                count,
                round(part_area, 4),
                round(part_area * thickness, 4),
                round(sum(_hole_area(h) for h in result.get("holes", [])), 4),
                _format_weight(_weight_kg(part_area * count, thickness, density)),
                _weight_formula(part_area, thickness, density, count=count),
            ])
        for col_idx in range(1, parts_ws.max_column + 1):
            parts_ws.column_dimensions[get_column_letter(col_idx)].width = 22

        wb.save(dst_path)
        logger.info("Excel を保存しました: %s", dst_path)

    def append_to_summary(
        self,
        results: list[dict],
        candidate: dict,
        params: dict,
        output_dir: str | Path = "output",
        *,
        candidates: list[dict] | None = None,
    ) -> None:
        """Summary に追記し、部品内訳と全案の配置図を更新する。"""
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        summary_path = output_dir / "summary.xlsx"

        sheet_size = f"{params.get('sheet_w', 0):.0f}×{params.get('sheet_h', 0):.0f} mm"
        part_names = ", ".join(
            r.get("display_name") or Path(r["step_path"]).stem for r in results
        )
        count_per_part = candidate.get("count_per_part", {})
        thickness = params.get("thickness", 0)
        density = params.get("density", 0)
        total_used_area = candidate.get("total_used_area", 0)
        scrap_area = _scrap_area(candidate)
        new_row = [
            datetime.now(tz=UTC).astimezone().strftime("%Y-%m-%d %H:%M:%S"),
            part_names,
            candidate.get("strategy_name", ""),
            sum(count_per_part.values()),
            round(candidate.get("yield_pct", 0), 2),
            round(candidate.get("remaining_area", 0), 4),
            round(candidate.get("total_hole_area", 0), 4),
            sheet_size,
            density,
            round(_weight_kg(total_used_area, thickness, density), 4),
            round(_weight_kg(scrap_area, thickness, density), 4),
            _weight_formula(total_used_area, thickness, density),
            _weight_formula(scrap_area, thickness, density),
        ]

        if summary_path.exists():
            wb = openpyxl.load_workbook(summary_path)
            ws = wb["Summary"] if "Summary" in wb.sheetnames else (wb.active or wb.create_sheet())
        else:
            wb = openpyxl.Workbook()
            ws = wb.active or wb.create_sheet()
            ws.title = "Summary"
            # ヘッダー行
            header_font = Font(bold=True)
            for col_idx, header in enumerate(_SUMMARY_HEADERS, start=1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for col_idx in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = 22

        ws.append(new_row)
        if "部品内訳" in wb.sheetnames:
            del wb["部品内訳"]
        detail_ws = wb.create_sheet("部品内訳")
        for label, area in [("部材重量合計", total_used_area), ("端材重量", scrap_area)]:
            weight = _weight_kg(area, thickness, density)
            detail_ws.append([
                label,
                f"{weight:.3f}kg ({floor(weight * 1000 + 0.5)}g)",
                _weight_formula(area, thickness, density),
            ])
        detail_ws.append([])
        detail_ws.append([
            "部品名", "数量指定", "配置数", "横(mm)", "縦(mm)",
            "部品面積(mm²)", "部品体積(mm³)", "重量", "重量 計算式",
        ])
        for index, result in enumerate(results):
            part_id = result.get("part_id") or f"{Path(result['step_path']).stem}_{index}"
            quantity = result.get("quantity", 1)
            count = count_per_part.get(part_id, 0)
            area = result.get("flattened_area", 0)
            weight = _weight_kg(area * count, thickness, density)
            detail_ws.append([
                result.get("display_name") or Path(result["step_path"]).stem,
                "自動(最大)" if quantity <= 0 else quantity,
                count,
                result.get("width", 0),
                result.get("height", 0),
                area,
                area * thickness,
                f"{weight:.3f}kg ({floor(weight * 1000 + 0.5)}g)",
                _weight_formula(area, thickness, density, count=count),
            ])
        for cell in detail_ws[4]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in detail_ws.iter_rows(min_row=5, min_col=4, max_col=7):
            for cell in row:
                cell.number_format = "0.00"
        for column, width in [("A", 40), ("B", 24), ("C", 18), ("D", 14), ("E", 14),
                              ("F", 22), ("G", 22), ("H", 24), ("I", 85)]:
            detail_ws.column_dimensions[column].width = width
        for row_index in (1, 2):
            detail_ws.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=9)
        detail_ws.freeze_panes = "D5"

        for name in wb.sheetnames:
            if name == "配置図" or (name.startswith("配置図_案") and name.removeprefix("配置図_案").isdigit()):
                del wb[name]
        layout_candidates = candidates or [candidate]
        for index, layout_candidate in enumerate(layout_candidates, start=1):
            title = f"配置図_案{index}" if len(layout_candidates) > 1 else "配置図"
            layout_ws = wb.create_sheet(title)
            layout_ws["A1"] = (
                f"配置図 (案{index}: {layout_candidate.get('strategy_name', '')}, "
                f"歩留まり {round(layout_candidate.get('yield_pct', 0), 1)}%)"
            )
            layout_ws["A1"].font = Font(bold=True)
            try:
                png_bytes = render_nesting_png(
                    results, layout_candidate, params.get("sheet_w", 0), params.get("sheet_h", 0),
                )
                image = XLImage(BytesIO(png_bytes))
                image.width = max(1, round(image.width / 2))
                image.height = max(1, round(image.height / 2))
                layout_ws.add_image(image, "A3")
            except Exception:
                logger.warning("配置図の画像埋め込みに失敗しました (案%s)", index, exc_info=True)

        wb.save(summary_path)
        logger.info("summary.xlsx に行を追記しました: %s", summary_path)


def _hole_area(points: list[tuple[float, float]]) -> float:
    """穴ポリゴンの靴紐公式面積 (絶対値) を返す。"""
    n = len(points)
    if n < 3:
        return 0.0
    s = sum(
        points[i][0] * points[(i + 1) % n][1] - points[(i + 1) % n][0] * points[i][1]
        for i in range(n)
    )
    return abs(s) / 2.0
